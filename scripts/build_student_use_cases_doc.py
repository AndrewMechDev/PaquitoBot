"""Create the one-page student-facing PaquitoBot use case guide."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


root = Path(__file__).resolve().parents[1]
output = root / "requerimientos" / "CASOS_DE_USO_ESTADO_ACTUAL_2026-10-20.docx"

doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(.72)
section.bottom_margin = Inches(.72)
section.left_margin = Inches(.9)
section.right_margin = Inches(.9)

for name in ("Normal", "Title", "Heading 1"):
    style = doc.styles[name]
    style.font.name = "Arial"
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn("w:" + attr), "Arial")

normal = doc.styles["Normal"]
normal.font.size = Pt(11)
normal.paragraph_format.line_spacing = 1.12
normal.paragraph_format.space_after = Pt(6)

title = doc.styles["Title"]
title.font.size = Pt(20)
title.font.bold = True
title.paragraph_format.space_after = Pt(7)
title_ppr = title._element.get_or_add_pPr()
for border in title_ppr.findall(qn("w:pBdr")):
    title_ppr.remove(border)

heading = doc.styles["Heading 1"]
heading.font.size = Pt(12)
heading.font.bold = True
heading.paragraph_format.space_before = Pt(12)
heading.paragraph_format.space_after = Pt(5)
heading.paragraph_format.keep_with_next = True

doc.add_paragraph("Qué puedes hacer en PaquitoBot", "Title")
doc.add_paragraph(
    "PaquitoBot reúne tus actividades de estudio en un solo lugar. Estas son las opciones que encontrarás en la app."
)


def add_action(title_text: str, detail: str) -> None:
    p = doc.add_paragraph(style="Normal")
    p.style = doc.styles["List Bullet"]
    p.paragraph_format.left_indent = Inches(.24)
    p.paragraph_format.first_line_indent = Inches(-.14)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_together = True
    p.add_run(title_text + ". ").bold = True
    p.add_run(detail)


doc.add_heading("Para empezar", 1)
add_action("Entrar a tu cuenta", "Regístrate o inicia sesión con tu correo institucional para acceder a la app.")
add_action("Elegir tus avisos", "Selecciona qué recordatorios quieres recibir y cuándo prefieres recibirlos. Puedes cambiarlo después.")

doc.add_heading("Para organizar tus estudios", 1)
add_action("Ver tus pendientes", "En Inicio encontrarás las entregas próximas. Toca un día de la semana para ver solo las de esa fecha.")
add_action("Revisar tus cursos", "Consulta tus notas registradas, faltas y próximas entregas; abre un curso para ver más detalle.")
add_action("Mirar tu horario", "Consulta clases, laboratorios y fechas de entrega ordenados por día.")
add_action("Preguntarle a Paquito", "Haz preguntas sobre tus cursos, tareas, notas o asistencia y recibe una respuesta en el chat.")
add_action("Consultar tus notificaciones", "Lee tus avisos, filtra los recientes y elimina los que ya no necesites.")
add_action("Enviar tareas a Google Calendar", "Si la conexión con Google está disponible, elige qué entregas quieres añadir a tu calendario.")

doc.add_heading("Para conocer otros servicios", 1)
add_action("Explorar eventos", "Revisa actividades de interés y abre sus páginas para conocer los detalles.")
add_action("Consultar matrícula y trámites", "Sigue la guía de matrícula o busca los requisitos y pasos de un trámite. La guía de matrícula también se puede abrir antes de iniciar sesión.")
add_action("Ver tu perfil", "Consulta tu nombre y correo, o cierra sesión cuando termines.")

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(12)
p.paragraph_format.space_after = Pt(0)
p.add_run("Ten en cuenta: ").bold = True
p.add_run("por ahora, los cursos, notas, asistencias y tareas que muestra la app son ejemplos; no corresponden a tus registros reales de Tecsup.")

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run("PaquitoBot  |  Guía para estudiantes")
for run in footer.runs:
    run.font.name = "Arial"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(95, 95, 95)

doc.save(output)
print(output)
