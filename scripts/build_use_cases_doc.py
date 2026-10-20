"""Build the current Android use case catalogue for PaquitoBot."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "requerimientos" / "CASOS_DE_USO_ESTADO_ACTUAL_2026-10-20.docx"


CASES = [
    ("CU 01", "Completar la configuración inicial", "Acceso", "Operativo en Android",
     "Estudiante que abre la app por primera vez", "Instalación nueva o configuración inicial sin completar.",
     "Abre la app, revisa bienvenida, elige categorías de avisos, anticipación y resumen diario, decide sobre el permiso del sistema y termina el recorrido.",
     "Puede denegar el permiso; la configuración se guarda, pero los avisos del sistema requieren autorización. Puede volver entre los pasos.",
     "Quedan guardadas la configuración inicial y las preferencias de notificación."),
    ("CU 02", "Registrarse o iniciar sesión con Clerk", "Acceso", "Operativo con servicio externo",
     "Estudiante TECSUP", "Configuración inicial terminada y Clerk disponible.",
     "Abre el formulario de Clerk, se registra o inicia sesión y completa la verificación que solicite el proveedor. La app valida el correo principal verificado @tecsup.edu.pe antes de abrir los datos académicos.",
     "Un correo de otro dominio o no verificado queda bloqueado; la pantalla permite cerrar esa sesión y usar otro correo. Si Clerk tarda en cargar, se muestra un mensaje de espera y reintento.",
     "Se establece una sesión autorizada asociada al ID estable de Clerk."),
    ("CU 03", "Consultar perfil y cerrar sesión", "Acceso", "Operativo en Android",
     "Estudiante autenticado", "Sesión institucional válida.",
     "Abre Perfil, consulta nombre y correo de Clerk, pulsa Cerrar sesión y confirma la acción.",
     "Si Clerk falla al cerrar la sesión, la app informa el error y permite reintentar.",
     "Se retira el acceso a las pantallas protegidas; el siguiente ingreso requiere autenticación."),
    ("CU 04", "Asignar un escenario académico a la cuenta", "Datos de demostración", "Operativo con datos locales",
     "Sistema; beneficiario: estudiante autenticado", "Existe un ID de usuario Clerk autorizado.",
     "La app calcula y conserva la asignación de uno de 15 archivos JSON, crea el repositorio de la cuenta activa y entrega esos datos a Inicio, Cursos, Horarios, avisos, Calendar y chat.",
     "Dos cuentas pueden recibir el mismo escenario. Una falla al leer el JSON se comunica en las pantallas académicas con opción de reintentar.",
     "La cuenta vuelve a ver el mismo escenario mientras se mantenga el catálogo; los datos siguen siendo ficticios."),
    ("CU 05", "Consultar el resumen académico de Inicio", "Seguimiento académico", "Operativo con datos locales",
     "Estudiante autenticado", "Escenario académico cargado.",
     "Abre Inicio y consulta tareas pendientes, avisos destacados de faltas o notas bajas, y calendario de la semana actual.",
     "Si falla la carga, aparece un mensaje con Reintentar. Los indicadores de ranking y ciclo en Inicio son datos fijos de demostración, independientes del escenario asignado.",
     "El estudiante obtiene una vista resumida de su escenario."),
    ("CU 06", "Filtrar tareas de Inicio por día", "Seguimiento académico", "Operativo con datos locales",
     "Estudiante autenticado", "La vista Inicio muestra la semana actual y tareas sin nota con fecha vigente.",
     "Selecciona un día de la semana; la lista se limita a tareas cuya fecha coincide. Pulsa Ver todas para recuperar la lista completa.",
     "Si ese día no tiene tareas pendientes, se muestra el estado vacío. Las tareas calificadas, vencidas o sin fecha no cuentan como pendientes.",
     "La selección modifica solo la vista; no cambia el archivo JSON."),
    ("CU 07", "Consultar cursos y estado general", "Seguimiento académico", "Operativo con datos locales",
     "Estudiante autenticado", "Escenario académico cargado.",
     "Abre Cursos y revisa cada curso, promedio simple de notas registradas, faltas, límite de referencia y siguiente entrega pendiente.",
     "Un curso sin notas muestra un guion. Si falla la carga, se ofrece Reintentar. El límite de ocho faltas es una regla de demostración, no un dato institucional del JSON.",
     "El estudiante identifica cursos y entregas que requieren atención."),
    ("CU 08", "Consultar detalle de curso", "Seguimiento académico", "Operativo con datos locales",
     "Estudiante autenticado", "El curso está disponible en la lista.",
     "Selecciona un curso, despliega asistencia y entregas, y abre el detalle de una evaluación para ver fecha, estado y nota si existe.",
     "Las entregas sin calificación pueden aparecer pendientes o vencidas. La pantalla no calcula un simulador de notas ni distingue prácticas y laboratorios en secciones separadas.",
     "Quedan visibles los registros académicos del curso del escenario asignado."),
    ("CU 09", "Consultar horarios y actividades", "Seguimiento académico", "Operativo con datos locales",
     "Estudiante autenticado", "Escenario académico cargado.",
     "Abre Horarios y revisa por fecha las clases, laboratorios, entregas y sesiones con falta registrada de todos sus cursos.",
     "Si no hay actividades o falla la lectura, la pantalla muestra el estado correspondiente y permite reintentar cuando hay error.",
     "El estudiante obtiene una línea de tiempo académica consolidada."),
    ("CU 10", "Consultar al asistente académico", "Asistente IA", "Implementado; depende de MiniMax",
     "Estudiante autenticado", "Sesión Clerk válida, conexión y clave MiniMax configurada en la compilación.",
     "Abre el chat, escribe una pregunta académica y envía. La app construye un JSON de cursos, tareas, notas y asistencias del escenario de esa cuenta y lo adjunta a la solicitud a MiniMax.",
     "Si cambia la sesión, falta la clave o falla la API, se muestra un mensaje de error. La instrucción del asistente limita el tema a información académica del sistema.",
     "La respuesta aparece en la conversación de la sesión actual; no representa datos reales de Tecsup."),
    ("CU 11", "Configurar avisos académicos", "Notificaciones", "Operativo en Android",
     "Estudiante autenticado", "Acceso a la bandeja de avisos o al primer uso.",
     "Abre Preferencias, elige categorías, horas de anticipación y resumen diario; guarda los cambios y decide sobre el permiso de notificaciones de Android.",
     "Puede cancelar la edición desde la bandeja. Si el permiso del sistema está desactivado, no se programan avisos visibles.",
     "Las preferencias locales se actualizan y el programador de avisos se reajusta."),
    ("CU 12", "Recibir un aviso académico local", "Notificaciones", "Implementado con datos locales",
     "Sistema Android; beneficiario: estudiante autenticado", "Configuración inicial completa, categoría habilitada y notificaciones permitidas.",
     "El trabajo periódico consulta el escenario activo, detecta entregas próximas, notas de laboratorio, faltas recientes y resumen diario según las preferencias, y publica los avisos elegibles.",
     "No se avisa por una entrega ya calificada; se evita repetir avisos enviados. La periodicidad de WorkManager y los permisos del dispositivo pueden retrasar o impedir la entrega.",
     "El aviso aparece en Android y se guarda en el historial local de esa cuenta."),
    ("CU 13", "Gestionar la bandeja de avisos", "Notificaciones", "Operativo en Android",
     "Estudiante autenticado", "Hay sesión autorizada; puede haber o no avisos guardados.",
     "Abre Notificaciones, filtra por Todos, Hoy, Semana o Anteriores, consulta un aviso, cambia vista completa o compacta y elimina uno o todos los elementos.",
     "Puede deshacer la eliminación individual desde la barra temporal. Si no hay avisos, se muestra una pantalla vacía.",
     "El historial local queda actualizado para el usuario activo."),
    ("CU 14", "Buscar y consultar trámites de secretaría", "Servicios al estudiante", "Guía informativa local",
     "Estudiante autenticado", "Acceso desde Perfil; catálogo local de trámites disponible.",
     "Abre Secretaría, busca por nombre o código, filtra por categoría y entra al detalle para revisar requisitos, pasos, costo referencial, plazo y canales externos.",
     "Si no hay coincidencias, puede limpiar los filtros; si falla el catálogo, puede reintentar. Los costos y condiciones deben confirmarse con Tecsup.",
     "El estudiante recibe orientación; la app no radica ni paga el trámite."),
    ("CU 15", "Recorrer la guía de matrícula", "Servicios al estudiante", "Guía informativa local",
     "Estudiante, incluso sin iniciar sesión", "Catálogo local de matrícula disponible.",
     "Abre la guía desde Bienvenida o Perfil, elige una ruta, revisa lista previa, pasos, canales y fechas de referencia; puede avanzar, retroceder o consultar un trámite relacionado.",
     "Cuando no hay cronograma confirmado, la guía lo indica y ofrece un canal oficial. Finalizar la lectura no acredita matrícula ni pago.",
     "El estudiante conoce el procedimiento y los canales para completarlo fuera de la app."),
    ("CU 16", "Explorar eventos", "Contenido informativo", "Muestra estática",
     "Estudiante autenticado", "Acceso desde Inicio.",
     "Abre Eventos, recorre la selección local, consulta la ficha de cada evento y puede abrir el enlace externo de inscripción o fuente.",
     "Las fichas son una muestra fechada y pueden perder vigencia; la inscripción se gestiona en el sitio externo.",
     "El estudiante obtiene información y acceso a la fuente del evento."),
    ("CU 17", "Autorizar Google Calendar", "Integración externa", "Implementado; validación real pendiente",
     "Estudiante autenticado", "Sesión Clerk válida y configuración OAuth de Google correcta para el paquete y firma instalados.",
     "Desde Horarios abre Google Calendar, pulsa Conectar, concede el permiso solicitado y revisa el correo de la cuenta Google conectada.",
     "Si cancela o falla autorización, permisos, configuración o red, la app muestra el motivo disponible y no envía tareas.",
     "Queda un token de corta duración en memoria, listo para sincronización manual."),
    ("CU 18", "Seleccionar y sincronizar tareas con Google Calendar", "Integración externa", "Implementado; validación real pendiente",
     "Estudiante autenticado", "Google autorizado, tareas sin nota con fecha válida y al menos una seleccionada.",
     "Revisa las tareas futuras seleccionadas por defecto, activa vencidas si quiere incluirlas, ajusta la selección y pulsa Sincronizar. La app crea o actualiza eventos propios en el calendario principal.",
     "Las tareas sin fecha válida se omiten. Un resultado parcial muestra creados, actualizados, sin cambios y fallidos; se puede reintentar sin duplicar eventos identificados por la integración.",
     "Se informa el resultado del envío. La operación es manual y de una sola dirección."),
    ("CU 19", "Revocar acceso de Google Calendar", "Integración externa", "Implementado; validación real pendiente",
     "Estudiante autenticado", "Cuenta Google conectada en la pantalla actual.",
     "Pulsa Revocar acceso; la app solicita revocación a Google y descarta credenciales en memoria.",
     "Si Google no responde, se muestra error y puede reintentar. Los eventos enviados previamente permanecen en Calendar.",
     "La siguiente sincronización requiere volver a autorizar Google."),
]


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def borders(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    bd = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "D9D9D9")
        bd.append(el)
    tc_pr.append(bd)


def cell_padding(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for side, value in (("top", 90), ("bottom", 90), ("left", 100), ("right", 100)):
        el = OxmlElement("w:" + side)
        el.set(qn("w:w"), str(value))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tc_pr.append(mar)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(.72)
section.bottom_margin = Inches(.7)
section.left_margin = Inches(.82)
section.right_margin = Inches(.82)

styles = doc.styles
styles["Normal"].font.name = "Arial"
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].font.color.rgb = RGBColor(20, 20, 20)
styles["Normal"].paragraph_format.space_after = Pt(3)
styles["Normal"].paragraph_format.line_spacing = 1.06
for name, size, before, after in (("Title", 20, 0, 9), ("Heading 1", 14, 14, 6), ("Heading 2", 11.5, 10, 4)):
    st = styles[name]
    st.font.name = "Arial"
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True
for name in ("Normal", "Title", "Heading 1", "Heading 2"):
    rpr = styles[name]._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn("w:" + attr), "Arial")
title_style_ppr = styles["Title"]._element.get_or_add_pPr()
for old_border in title_style_ppr.findall(qn("w:pBdr")):
    title_style_ppr.remove(old_border)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run("PaquitoBot  ·  Casos de uso  ·  20 octubre 2026  |  ")
fld = OxmlElement("w:fldSimple")
fld.set(qn("w:instr"), "PAGE")
footer._p.append(fld)
footer.style = styles["Normal"]
for run in footer.runs:
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(90, 90, 90)

doc.add_paragraph("Casos de uso del estado actual de PaquitoBot", "Title")
p = doc.add_paragraph()
p.add_run("Alcance. ").bold = True
p.add_run("Aplicación Android del repositorio al 20 de octubre de 2026. Este catálogo describe lo que permite el código actual y sirve como base para documentación funcional, revisión del MVP y futuras pruebas de aceptación.")
p = doc.add_paragraph()
p.add_run("Lectura clave. ").bold = True
p.add_run("Clerk autentica la identidad institucional. Los cursos, notas, asistencias y tareas proceden de 15 escenarios JSON de demostración asignados a la cuenta; no son registros reales de Tecsup. Google Calendar está implementado en código, pero su flujo real de OAuth y escritura aún necesita validación en dispositivo.")

doc.add_heading("Actores y límites", 1)
p = doc.add_paragraph()
p.add_run("Actor principal: ").bold = True
p.add_run("estudiante con correo principal @tecsup.edu.pe verificado en Clerk. La guía de matrícula también está disponible antes del inicio de sesión.")
p = doc.add_paragraph()
p.add_run("Sistemas de apoyo: ").bold = True
p.add_run("Clerk para la sesión; MiniMax para respuestas del chat; Android y WorkManager para avisos; Google Identity Services y Google Calendar para exportar tareas.")
p = doc.add_paragraph()
p.add_run("Plataformas: ").bold = True
p.add_run("los casos siguientes describen Android. El proyecto contiene un punto de entrada iOS, pero su ContentView sigue siendo una pantalla base sin estos flujos.")

doc.add_heading("Índice de casos", 1)
table = doc.add_table(rows=1, cols=3)
table.autofit = False
for cell, width in zip(table.rows[0].cells, (Inches(.72), Inches(4.23), Inches(1.88))):
    cell.width = width
for i, label in enumerate(("ID", "Caso", "Área")):
    table.rows[0].cells[i].text = label
    shade(table.rows[0].cells[i], "253B58")
    for run in table.rows[0].cells[i].paragraphs[0].runs:
        run.font.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        run.font.size = Pt(9)
for n, case in enumerate(CASES):
    row = table.add_row()
    for i, value in enumerate((case[0], case[1], case[2])):
        row.cells[i].text = value
        row.cells[i].width = (Inches(.72), Inches(4.23), Inches(1.88))[i]
        if n % 2:
            shade(row.cells[i], "F3F6F9")
for row in table.rows:
    for cell in row.cells:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        borders(cell)
        cell_padding(cell)
        for p in cell.paragraphs:
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.03
            for run in p.runs:
                run.font.size = Pt(8.5)

doc.add_heading("Descripción de los casos", 1)
current_area = None
for cid, title, area, status, actor, pre, flow, alt, result in CASES:
    if area != current_area:
        doc.add_heading(area, 2)
        current_area = area
    h = doc.add_paragraph(style="Heading 2")
    h.add_run(f"{cid}  {title}")
    meta = doc.add_paragraph()
    meta.paragraph_format.space_after = Pt(5)
    meta.paragraph_format.keep_with_next = True
    meta.add_run("Actor  ").bold = True
    meta.add_run(actor)
    meta.add_run("     Estado  ").bold = True
    meta.add_run(status)
    for label, value in (("Precondición", pre), ("Flujo principal", flow), ("Alternativas y excepciones", alt), ("Resultado", result)):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(.15)
        p.paragraph_format.keep_with_next = label != "Resultado"
        p.add_run(label + ". ").bold = True
        p.add_run(value)
    doc.paragraphs[-1].paragraph_format.space_after = Pt(7)

doc.add_heading("Notas para la documentación funcional", 1)
for lead, detail in (
    ("Fuente de datos académicos", "academic/catalog.json y academic/scenarios/scenario-01.json a scenario-15.json. La asignación se deriva del ID de Clerk y se guarda en preferencias privadas de la instalación. No garantiza escenarios exclusivos entre estudiantes."),
    ("Reglas de presentación", "Las tareas pendientes de Inicio y Cursos son entregas sin nota con fecha de hoy o futura; el detalle de un curso también permite revisar entregas vencidas y calificadas. El promedio mostrado es la media de notas registradas y el límite de faltas es de referencia."),
    ("Servicios externos", "El chat envía el JSON académico de la cuenta activa a MiniMax desde el cliente Android. Calendar exporta tareas manualmente a Google y no importa cambios. Ambos flujos requieren credenciales, red y configuración externa."),
    ("Fuera del alcance actual", "No hay conexión con el sistema académico real de Tecsup, simulador de notas, tramitación o pago dentro de la app, inscripción a eventos dentro de la app ni equivalencia funcional en iOS."),
):
    p = doc.add_paragraph()
    p.add_run(lead + ". ").bold = True
    p.add_run(detail)

doc.add_heading("Trazabilidad en el repositorio", 1)
for label, path in (
    ("Navegación y acceso", "androidApp/src/main/kotlin/pe/tecsup/paquitobot/MainActivity.kt; session/SessionViewModel.kt; ui/auth/ClerkAuthGate.kt"),
    ("Escenarios y vistas académicas", "androidApp/src/main/assets/academic/; PaquitoBotApp.kt; ui/academic/; ui/home/; ui/courses/; ui/schedule/"),
    ("IA y avisos", "data/AcademicContextSnapshot.kt; data/MiniMaxChatRepository.kt; data/notifications/; ui/notifications/"),
    ("Guías, eventos y calendario", "assets/procedures.json; ui/procedures/; ui/events/; ui/calendar/; data/calendar/"),
):
    p = doc.add_paragraph()
    p.add_run(label + ". ").bold = True
    p.add_run(path)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
