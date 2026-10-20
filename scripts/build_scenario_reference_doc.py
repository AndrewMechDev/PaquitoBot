"""Generate a source-derived Word reference for the 15 academic demo scenarios."""

import json
from collections import defaultdict
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "androidApp" / "src" / "main" / "assets" / "academic"
OUTPUT = ROOT / "requerimientos" / "GUIA_15_ESCENARIOS_ACADEMICOS.docx"
AS_OF = date(2026, 10, 20)
ABSENCE_LIMIT = 8  # El mismo valor de referencia mostrado por la app.
GRADE_COLUMNS = (
    ("L1", "Laboratorio 1"), ("L2", "Laboratorio 2"),
    ("L3", "Laboratorio 3"), ("L4", "Laboratorio 4"),
    ("P1", "Practica 1"), ("P2", "Practica 2"), ("P3", "Practica 3"),
    ("F1", "Foro 1"), ("F2", "Foro 2"),
    ("D", "Diseño de red"), ("C", "Cuestionario de protocolos"),
    ("T", "Proyecto de topología"),
)


def read_json(path):
    return json.loads(path.read_text(encoding="utf-8"))


catalog = read_json(DATA / "catalog.json")
scenarios = [read_json(p) for p in sorted((DATA / "scenarios").glob("scenario-*.json"))]
assignments = {item["id"]: item for item in catalog["assignments"]}
sessions = {item["id"]: item for item in catalog["class_sessions"]}


def score(value):
    return str(int(value)) if float(value).is_integer() else f"{value:.1f}".replace(".", ",")


def average(values):
    return "—" if not values else f"{sum(values) / len(values):.1f}".replace(".", ",")


def short_date(value):
    return f"{value[8:10]}/{value[5:7]}"


def normalized_course(scenario, course):
    cid = course["id"]
    grades = {item["assignment_id"]: item["score"] for item in scenario["grades"]}
    course_assignments = [item for item in catalog["assignments"] if item["course_id"] == cid]
    own_grades = [grades[item["id"]] for item in course_assignments if item["id"] in grades]
    own_attendance = [item for item in scenario["attendance"] if sessions[item["class_session_id"]]["course_id"] == cid]
    absent = [item for item in own_attendance if item["status"] == "absent"]
    lab_absent = [item for item in absent if sessions[item["class_session_id"]].get("kind") == "laboratory"]
    absent_dates = sorted(
        (sessions[item["class_session_id"]].get("start_at") or "")[:10]
        for item in absent
        if sessions[item["class_session_id"]].get("start_at")
    )
    overrides = scenario.get("due_at_overrides", {})
    pending, overdue, no_date = [], [], []
    for item in course_assignments:
        if item["id"] in grades:
            continue
        due = overrides.get(str(item["id"]), item.get("due_at"))
        if not due:
            no_date.append(item)
        elif due[:10] < AS_OF.isoformat():
            overdue.append((item, due))
        else:
            pending.append((item, due))
    pending.sort(key=lambda pair: pair[1])
    return {
        "course": course,
        "average": average(own_grades),
        "grade_count": len(own_grades),
        "absences": len(absent),
        "lab_absences": len(lab_absent),
        "absent_dates": absent_dates,
        "pending": pending,
        "overdue": overdue,
        "no_date": no_date,
        "grades_by_name": {item["name"]: grades.get(item["id"]) for item in course_assignments},
    }


records = [(scenario, [normalized_course(scenario, course) for course in catalog["courses"]]) for scenario in scenarios]


def shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    old = tcpr.find(qn("w:shd"))
    if old is not None:
        tcpr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def border_and_padding(cell):
    tcpr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        node = OxmlElement("w:" + edge)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), "D9D9D9")
        borders.append(node)
    tcpr.append(borders)
    margins = OxmlElement("w:tcMar")
    for side, value in (("top", 105), ("bottom", 105), ("left", 105), ("right", 105)):
        node = OxmlElement("w:" + side)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    tcpr.append(margins)


def style_table(table, widths, font_size=8.5, header_rows=1):
    table.autofit = False
    for index, width in enumerate(widths):
        table.columns[index].width = Inches(width)
    for row_number, row in enumerate(table.rows):
        for index, cell in enumerate(row.cells):
            cell.width = Inches(widths[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            border_and_padding(cell)
            if row_number < header_rows:
                shade(cell, "233A54" if row_number == 0 else "3C5874")
            elif (row_number - header_rows) % 2 == 1:
                shade(cell, "F2F6F8")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.03
                if index > 0:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    run.font.size = Pt(font_size)
                    if row_number < header_rows:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(11)
section.page_height = Inches(8.5)
section.top_margin = Inches(.55)
section.bottom_margin = Inches(.52)
section.left_margin = Inches(.6)
section.right_margin = Inches(.6)

for name in ("Normal", "Title", "Heading 1", "Heading 2"):
    style = doc.styles[name]
    style.font.name = "Arial"
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn("w:" + attr), "Arial")
normal = doc.styles["Normal"]
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.1
for name, size, before, after in (("Title", 20, 0, 9), ("Heading 1", 14, 13, 6), ("Heading 2", 11, 9, 4)):
    style = doc.styles[name]
    style.font.size = Pt(size)
    style.font.bold = True
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True
title_ppr = doc.styles["Title"]._element.get_or_add_pPr()
for old in title_ppr.findall(qn("w:pBdr")):
    title_ppr.remove(old)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run("PaquitoBot  |  Escenarios de demostración  |  ")
page = OxmlElement("w:fldSimple")
page.set(qn("w:instr"), "PAGE")
footer._p.append(page)
for run in footer.runs:
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(100, 100, 100)

doc.add_paragraph("Guía de los 15 escenarios académicos", "Title")
doc.add_paragraph("Busca en la esquina superior izquierda de Inicio el ID «Escenario XX» de tu cuenta. Luego compara los datos de esta guía con Cursos, el detalle de cada curso e Inicio.")
doc.add_paragraph("Los datos son de demostración. Los promedios son la media simple de las notas registradas; las faltas se comparan con un límite de referencia de 8 por curso. Las entregas pendientes se calcularon al 20/10/2026 y cambian al pasar los días.")

overview = doc.add_table(rows=1, cols=6)
for cell, label in zip(overview.rows[0].cells, ("ID", "Situación", "Notas", "Promedio", "Faltas", "Pendientes")):
    cell.text = label
for scenario, courses in records:
    row = overview.add_row()
    all_grades = [item["score"] for item in scenario["grades"]]
    values = (scenario["scenario_id"], scenario["label"], str(len(all_grades)), average(all_grades),
              str(sum(course["absences"] for course in courses)), str(sum(len(course["pending"]) for course in courses)))
    for cell, value in zip(row.cells, values):
        cell.text = value
style_table(overview, (.65, 4.25, .95, 1.2, 1.25, 1.6), font_size=9.2)

for scenario, courses in records:
    doc.add_page_break()
    doc.add_paragraph(f"Escenario {scenario['scenario_id']}  {scenario['label']}", "Title")
    total_grades = len(scenario["grades"])
    all_scores = [grade["score"] for grade in scenario["grades"]]
    total_absences = sum(item["absences"] for item in courses)
    total_lab_absences = sum(item["lab_absences"] for item in courses)
    total_pending = sum(len(item["pending"]) for item in courses)
    total_overdue = sum(len(item["overdue"]) for item in courses)
    p = doc.add_paragraph()
    p.add_run("Resumen  ").bold = True
    p.add_run(f"{total_grades} notas registradas · promedio {average(all_scores)}/20 · {total_absences} faltas, {total_lab_absences} en laboratorios · {total_pending} entregas pendientes · {total_overdue} vencidas sin nota.")

    doc.add_heading("Lo que debes ver por curso", 1)
    table = doc.add_table(rows=2, cols=18)
    top_labels = {0: "Curso", 1: "Prom.", 2: "Laboratorios", 6: "Prácticas",
                  9: "Foros", 11: "Otras notas", 14: "Faltas",
                  15: "En lab", 16: "Pend.", 17: "Venc."}
    for index, label in top_labels.items():
        table.cell(0, index).text = label
    for index, (short_label, _) in enumerate(GRADE_COLUMNS, start=2):
        table.cell(1, index).text = short_label
    for item in courses:
        row = table.add_row()
        grade_values = []
        for _, name in GRADE_COLUMNS:
            if name not in item["grades_by_name"]:
                grade_values.append("—")
            else:
                value = item["grades_by_name"][name]
                grade_values.append(score(value) if value is not None else "s/n")
        values = (item["course"]["code"] + "\n" + item["course"]["name"], item["average"],
                  *grade_values, f"{item['absences']}/{ABSENCE_LIMIT}",
                  str(item["lab_absences"]), str(len(item["pending"])), str(len(item["overdue"])))
        for cell, value in zip(row.cells, values):
            cell.text = value
    style_table(table, (1.24, .5, *([.48] * 12), .62, .58, .5, .5),
                font_size=8.2, header_rows=2)
    for start, end in ((2, 5), (6, 8), (9, 10), (11, 13)):
        table.cell(0, start).merge(table.cell(0, end))
    for index in (0, 1, 14, 15, 16, 17):
        table.cell(0, index).merge(table.cell(1, index))

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(2)
    p.add_run("Notas sobre 20. ").bold = True
    p.add_run("L = laboratorio; P = práctica; F = foro; D = diseño de red; C = cuestionario de protocolos; T = proyecto de topología. “s/n” = actividad sin nota; “—” = actividad no incluida en el curso.")

    doc.add_heading("Faltas registradas", 1)
    courses_with_absences = [item for item in courses if item["absent_dates"]]
    if not courses_with_absences:
        doc.add_paragraph("No hay faltas registradas en este escenario.")
    else:
        for item in courses_with_absences:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(.12)
            p.paragraph_format.space_after = Pt(3)
            p.add_run(item["course"]["code"] + "  ").bold = True
            p.add_run(", ".join(short_date(day) for day in item["absent_dates"]))

    overdue_items = [(item["course"]["code"], task["name"], due)
                     for item in courses for task, due in item["overdue"]]
    if overdue_items:
        doc.add_heading("Entregas vencidas sin nota", 1)
        doc.add_paragraph("; ".join(f"{code} · {name} ({short_date(due)})" for code, name, due in overdue_items))

    doc.add_heading("Cómo comprobarlo en la app", 1)
    doc.add_paragraph("En Cursos, compara el promedio y las faltas de cada fila. Abre el curso para revisar las notas de cada entrega y el historial de asistencia. En Inicio, compara las entregas aún pendientes; el conteo puede bajar si la fecha ya pasó.")

    print(f"{scenario['scenario_id']} | {scenario['label']} | promedio {average(all_scores)} | faltas {total_absences} | laboratorios con falta {total_lab_absences} | pendientes {total_pending}")

doc.save(OUTPUT)
print(OUTPUT)
